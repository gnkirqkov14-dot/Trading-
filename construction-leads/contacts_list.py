"""Таблица за обаждания: фирми с телефон или имейл по градове (София, Пловдив) и етапът на сградите.

    python contacts_list.py
Изход: output/kontakti_sofia_plovdiv.xlsx

Само фирми с поне един контакт. Една фирма = един ред (фирми със същия телефон или имейл са
обединени – да не се звъни два пъти). Етапът е приблизителен: изчислен от датата на разрешението
и големината на сградата (Акт 14 не е публичен), освен където е потвърден с документ.
"""
from __future__ import annotations

import json
import re
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).parent
OUT = ROOT / "output" / "kontakti_sofia_plovdiv.xlsx"
TODAY = date.today()

STATUSES = ["Не е звъняно", "Звъннато – няма отговор", "Разговор – интерес", "Изпратена оферта",
            "Среща / оглед", "Не се интересува", "Грешен номер"]

# Етап -> (кратко описание, приоритет, ред за подреждане)
STAGE_TEXT = {
    "late": ("Около Акт 14 – фасадата се решава сега", "Висок – спешно", 0),
    "rough": ("Груб строеж – фасадата още не е поръчана", "Висок", 1),
    "hole": ("Изкоп / начало на строежа – рано, материалите още се избират", "Среден", 2),
    "reno": ("Предстои / тече обновяване на фасадата", "Среден", 3),
    "public": ("Обществена поръчка – строителството вероятно тече", "Среден", 4),
    "design": ("Проектиране – строежът предстои", "Среден (рано)", 5),
}


# ---------------------------------------------------------------------------
# Помощни
# ---------------------------------------------------------------------------

def fmt_phone(p: str) -> str:
    d = re.sub(r"\D", "", p or "")
    if d.startswith("359"):
        d = d[3:]
    d = d.lstrip("0")
    if len(d) == 9 and d[0] in "89":
        return f"+359 {d[:2]} {d[2:5]} {d[5:]}"
    if len(d) == 8 and d[0] == "2":
        return f"+359 2 {d[1:4]} {d[4:]}"
    if len(d) == 8:
        return f"+359 {d[:2]} {d[2:5]} {d[5:]}"
    return p


def month(iso: str) -> str:
    return f"{iso[5:7]}.{iso[:4]}" if iso and len(iso) >= 7 else ""


LOWER_WORDS = {"жк", "ж.к.", "кв.", "м.", "с.", "гр.", "ул.", "бул.", "в.з.", "и", "на", "в", "за", "от", "до", "с"}
UPPER_WORDS = re.compile(r"^(?:[IVX]+|нпз|жг|ЦГЧ|СГЧ|ЮИПЗ|ЮИЗ|СИЗ|ИИЗ|МФЗ|СМФ|ПИ|УПИ|ТП|ЕАД|ЕООД|ООД|АД|ЕТ)$", re.I)


def title(s: str) -> str:
    """Изречение с малки букви за описания, изписани само с главни („ЖИЛИЩНА СГРАДА“)."""
    s = re.sub(r"\s+", " ", s or "").strip(" ,;\"“”„")
    return s[:1].upper() + s[1:].lower() if s.isupper() else s


ROMAN = re.compile(r"^(?:I|II|III|IV|V|VI|VII|VIII|IX|X|XI|XII)$", re.I)
ACRONYMS = {"НПЗ", "ЖГ", "ЦГЧ", "СГЧ", "ЮИПЗ", "ЮИЗ", "СИЗ", "ИИЗ", "МФЗ", "СМФ", "ПИ", "УПИ", "ТП", "ЕАД", "ЕООД",
            "ООД", "АД", "ЕТ", "ЖР", "СО", "НДК", "ВЗ"}
ABBREV = {"м", "кв", "с", "гр", "ул", "бул", "жк", "ж", "к", "в", "з", "пл", "р-н", "местн"}


def place(s: str) -> str:
    """„М."ВОЛУЯК-ГАРАТА“ -> „м. „Волуяк-Гарата““ стил: всяка дума с главна, съкращения малки, римски цифри главни.
    Прилага се само за текст, изписан изцяло с главни букви."""
    s = re.sub(r"\s+", " ", s or "").strip(" ,;")
    if s.count('"') % 2:          # несдвоени кавички от регистъра („М.\"ВОЛУЯК“) – махаме ги
        s = s.replace('"', "")
    if s.count("„") != s.count("“") + s.count("”"):
        s = re.sub(r"[„“”]", "", s)
    if not s.isupper():
        return s

    def fix(m: re.Match) -> str:
        w = m.group(0)
        after = s[m.end():m.end() + 1]
        if (w.lower() in ABBREV and after == ".") or w.lower() in ("жк", "и", "на", "до", "от", "за", "в"):
            return w.lower()
        if ROMAN.match(w):
            return w.upper()
        if w.upper() in ACRONYMS:
            return w.upper()
        return w[:1].upper() + w[1:].lower()
    return re.sub(r"[A-Za-zА-Яа-я]+", fix, s)


def person(names) -> str:
    """Управители: „ИВАН ПЕТРОВ ИВАНОВ“ -> „Иван Петров Иванов“; без „ЕИК/ПИК …“ и „– Управител“."""
    out = []
    for n in names or []:
        n = re.split(r"(?i),\s*еик|\s+[–-]\s+управител", n)[0].strip(" ,\"”“")
        if not n:
            continue
        words = []
        for w in n.split():
            core = w.strip("„“”\"")
            if core.upper() in ACRONYMS:
                words.append(w.upper())
            elif w.isupper() or w.islower():
                words.append(re.sub(r"[A-Za-zА-Яа-я]+", lambda m: m.group(0)[:1].upper() + m.group(0)[1:].lower(), w, 1))
            else:
                words.append(w)
        out.append(" ".join(words))
    return ", ".join(dict.fromkeys(out))


ACCOUNTANT = re.compile(r"(?i)account|счетовод|schetovod|knigovod|accountant|tax|smetki|balans")


def mark_accountant(emails: list[str]) -> list[str]:
    """Имейлът в Търговския регистър често е на счетоводителя – отбелязваме го."""
    return [e + " (счетоводство)" if ACCOUNTANT.search(e) else e for e in emails]


def short(s: str, n: int = 70) -> str:
    s = re.sub(r"\s+", " ", s or "").strip()
    return s if len(s) <= n else s[:n - 1].rsplit(" ", 1)[0] + "…"


def name_of(c: dict) -> str:
    return re.sub(r"\s+", " ", c.get("name") or "").strip()


class Companies:
    """Фирми, обединени по ЕИК, телефон или имейл. Телефон/имейл, общ за повече от 6 фирми (счетоводна
    кантора, агенция за регистрация), не обединява – той не е на групата."""

    def __init__(self):
        self.rows: list[dict] = []

    def add(self, name: str, eik: str, role: str, phones, emails, website, person, obj: dict) -> None:
        phones = [p for p in dict.fromkeys(fmt_phone(p) for p in phones or []) if p]
        emails = [e for e in dict.fromkeys((e or "").strip().lower() for e in emails or []) if "@" in e]
        if not phones and not emails or not name:
            return
        self.rows.append({"names": [name], "eiks": [eik] if eik else [], "roles": [role], "phones": phones,
                          "emails": emails, "website": website or "", "person": person or "", "objects": [obj],
                          "id": f"eik:{eik}" if eik else f"name:{name.lower()}"})

    def merged(self) -> list[dict]:
        firms_per_key: dict[str, set] = {}
        for r in self.rows:
            for k in [f"ph:{p}" for p in r["phones"]] + [f"em:{e}" for e in r["emails"]]:
                firms_per_key.setdefault(k, set()).add(r["id"])
        parent = list(range(len(self.rows)))

        def find(i: int) -> int:
            while parent[i] != i:
                parent[i] = parent[parent[i]]
                i = parent[i]
            return i
        first: dict[str, int] = {}
        for i, r in enumerate(self.rows):
            keys = [r["id"]] + [k for k in [f"ph:{p}" for p in r["phones"]] + [f"em:{e}" for e in r["emails"]]
                                if len(firms_per_key[k]) <= 6]
            for k in keys:
                if k in first:
                    parent[find(i)] = find(first[k])
                else:
                    first[k] = i
        groups: dict[int, dict] = {}
        for i, r in enumerate(self.rows):
            g = groups.setdefault(find(i), {"names": [], "eiks": [], "roles": [], "phones": [], "emails": [],
                                            "website": "", "person": "", "objects": []})
            for k in ("names", "eiks", "roles", "phones", "emails"):
                g[k] += [x for x in r[k] if x not in g[k]]
            from leads.ksb import normalize_name
            seen, uniq = set(), []
            for n in g["names"]:
                if normalize_name(n) not in seen:
                    seen.add(normalize_name(n))
                    uniq.append(n)
            g["names"] = uniq
            g["website"] = g["website"] or r["website"]
            g["person"] = g["person"] or r["person"]
            g["objects"] += [o for o in r["objects"] if o["key"] not in {x["key"] for x in g["objects"]}]
        out = list(groups.values())
        for g in out:
            g["objects"].sort(key=lambda o: (STAGE_TEXT[o["code"]][2],
                                             -int(o["date"].replace("-", "")) if o["date"] else 0))
            g["best"] = g["objects"][0]
        out.sort(key=lambda g: (STAGE_TEXT[g["best"]["code"]][2], -len(g["objects"]), g["names"][0].lower()))
        return out


def stage_line(o: dict) -> str:
    text = STAGE_TEXT[o["code"]][0]
    extra = []
    if o.get("confirmed"):
        extra.append(o["confirmed"])
    if o.get("date"):
        extra.append(("договор от " if o["code"] in ("public", "design") else "разрешение от ") + month(o["date"]))
    return text + (f" ({'; '.join(extra)})" if extra else "")


# ---------------------------------------------------------------------------
# София
# ---------------------------------------------------------------------------

def sofia() -> tuple[list[dict], list[dict]]:
    global INFRA
    from plovdiv_permits import INFRA
    from pilot import load_raw
    from run import finalize
    permits = finalize(load_raw(), TODAY.isoformat())
    firms, supervision = Companies(), Companies()
    for p in permits:
        if not p.get("facade_window"):
            continue
        code = "reno" if p["kind"] == "Фасадна реновация" else p.get("stage_code")
        if code not in STAGE_TEXT or INFRA.search(p.get("object") or "") or \
                re.search(r"(?i)физическа инфраструктура|електронна съобщителна|климатичн|асансьор|рекламн|"
                          r"билборд|^монтаж", p.get("object") or ""):
            continue
        started = next((m for m in p.get("milestones") or [] if "започнал" in m.get("what", "")), None)
        where = ", ".join(x for x in [p.get("region"), place((p.get("locality") or "").replace('";"', ", "))] if x)
        obj = {"key": p.get("hash") or p["number"], "code": code, "date": p.get("in_force") or "",
               "where": where, "what": title(short(p.get("object"), 60)), "url": p.get("url", ""),
               "confirmed": "потвърдено начало на строежа" if started else "", "permit": p["number"]}
        if p.get("project_name"):
            obj["what"] = f'{p["project_name"]["name"]} – {obj["what"]}'
        for c in p.get("contacts", []):
            role = c.get("role", "")
            person_ = person((c.get("managers") or [])[:2])
            if role.startswith("Инвеститор"):
                firms.add(name_of(c), c.get("eik", ""), "Инвеститор (и строител по КСБ)" if "КСБ" in role else "Инвеститор",
                          c.get("phones"), c.get("emails"), c.get("website"), person_, obj)
            elif role == "Строител":
                firms.add(name_of(c), c.get("eik", ""), "Строител", c.get("phones"), c.get("emails"),
                          c.get("website"), person_, obj)
            elif role == "Архитект – потвърдена":
                firms.add(name_of(c), c.get("eik", ""), "Архитект", c.get("phones"), c.get("emails"),
                          c.get("website"), person_, obj)
            elif role == "Строителен надзор":
                supervision.add(name_of(c), c.get("eik", ""), "Строителен надзор", c.get("phones"), c.get("emails"),
                                c.get("website"), person_, obj)
    return firms.merged(), supervision.merged()


# ---------------------------------------------------------------------------
# Пловдив
# ---------------------------------------------------------------------------

def plovdiv() -> list[dict]:
    import run
    run.OFFLINE = True  # контактите вече са изтеглени – само от кеша
    import plovdiv_permits as P
    from leads.runlog import RunLog
    log = RunLog("contacts_list")
    objs = P.objects(P.read(P.RAW / "plovdiv_razreshenia.csv"), P.read(P.RAW / "plovdiv_vavedeni.csv"), "2023-01-01")
    comp = P.contacts(objs, log)
    firms = Companies()
    for o in objs:
        if not o["facade"]:
            continue
        code = "reno" if o["kind"] == "Фасадна реновация" else o["stage_code"]
        if code not in STAGE_TEXT:
            continue
        where = ", ".join(x for x in ["Пловдив", o["district"], place(o["part"])] if x)
        obj = {"key": f'plv-{o["permit"]}', "code": code, "date": o["date"], "where": where,
               "what": title(short(o["object"], 60)), "url": o["url"], "permit": o["permit"]}
        for name in o["companies"]:
            c = comp.get(name) or {}
            firms.add(name, c.get("eik", ""), "Инвеститор" + (" (и строител по КСБ)" if c.get("is_builder") else ""),
                      c.get("phones"), c.get("emails"), c.get("website"), person((c.get("managers") or [])[:2]), obj)
    # Обществени поръчки за сгради в областта (изпълнители и проектанти с контакт), последните 18 месеца
    cutoff = (TODAY - timedelta(days=548)).isoformat()
    public = json.loads((ROOT / "output" / "plovdiv" / "obekti.json").read_text(encoding="utf-8"))
    for o in public:
        if not o["facade_relevant"] or o["date"] < cutoff or o["contractor_role"] == "строителен надзор":
            continue
        head = re.sub(r'^[„"\s]+', "", o["object"]).split(":")[0].lower()
        if re.search(r"строителен надзор|оценка (?:за|на) съответстви", head) and \
                not re.search(r"проектиране|изработване на инвестиционн|авторски надзор", head):
            continue  # консултант/строителен надзор, не проектант или строител
        code = "design" if o["contractor_role"] == "проектант" else "public"
        if code == "design" and re.search(r"авторски надзор", head) and not re.search(r"проектиране|изработване", head):
            role_note = "Проектант (авторски надзор)"
            code = "public"  # авторски надзор = строежът тече
        else:
            role_note = ""
        role = {"design": "Проектант (обществена поръчка)"}.get(code, "Строител (обществена поръчка)")
        if "инженеринг" in o["contractor_role"]:
            role = "Проектиране и строителство (обществена поръчка)"
        role = role_note or role
        obj = {"key": f'eop-{o["procurement_no"]}-{o["contractor_eik"]}', "code": code, "date": o["date"],
               "where": o["place"] + (f' – възложител {place(short(o["buyer"], 45))}' if o.get("buyer") else ""),
               "what": title(short(re.sub(r'^[„"\s]+', "", o["object"]), 60)), "url": o["url"], "permit": ""}
        firms.add(o["contractor"], o.get("contractor_eik", ""), role, o.get("contractor_phones"),
                  o.get("contractor_emails"), o.get("contractor_website"), "", obj)
    return firms.merged()


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------

FONT = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
PRIORITY_FILL = {"Висок – спешно": "F8CBAD", "Висок": "FCE4D6", "Среден": "E2EFDA", "Среден (рано)": "EDEDED"}
THIN = Side(style="thin", color="D9D9D9")

COLUMNS = [("№", 5), ("Фирма", 34), ("Роля", 20), ("Телефон", 18), ("Имейл", 30), ("Сайт", 22),
           ("Лице за контакт (управител)", 24), ("Брой сгради", 8), ("Сгради (къде и какво)", 60),
           ("Етап сега (приблизително)", 44), ("Приоритет", 13), ("Линк към обекта", 16),
           ("Статус", 22), ("Дата на обаждане", 13), ("Бележки / следваща стъпка", 40)]
INPUT_COLS = (13, 14, 15)


def sheet(wb: Workbook, title_: str, rows: list[dict], note: str) -> None:
    ws = wb.create_sheet(title_)
    ws["A1"] = f"{title_} – фирми с телефон или имейл ({len(rows)})"
    ws["A1"].font = Font(name=FONT, bold=True, size=13)
    ws["A2"] = note
    ws["A2"].font = Font(name=FONT, italic=True, size=9, color="595959")
    hdr = 4
    for i, (h, w) in enumerate(COLUMNS, 1):
        c = ws.cell(row=hdr, column=i, value=h)
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = HEAD_FILL if i not in INPUT_COLS else PatternFill("solid", fgColor="BF8F00")
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    for n, g in enumerate(rows, 1):
        best = g["best"]
        label, priority, _ = STAGE_TEXT[best["code"]]
        places = []
        for o in g["objects"][:4]:
            places.append(f'• {o["where"]}: {o["what"]} – {STAGE_TEXT[o["code"]][0].split(" – ")[0].lower()}')
        if len(g["objects"]) > 4:
            places.append(f'… и още {len(g["objects"]) - 4}')
        values = [n, " / ".join(g["names"][:3]) + (f' (+{len(g["names"]) - 3})' if len(g["names"]) > 3 else ""),
                  ", ".join(dict.fromkeys(g["roles"])), "\n".join(g["phones"][:3]),
                  "\n".join(mark_accountant(g["emails"][:3])),
                  g["website"].replace("https://", "").replace("http://", "").rstrip("/"), g["person"],
                  len(g["objects"]), "\n".join(places), stage_line(best), priority,
                  "отвори" if best.get("url") else "", "Не е звъняно", None, None]
        r = hdr + n
        for i, v in enumerate(values, 1):
            c = ws.cell(row=r, column=i, value=v)
            c.font = Font(name=FONT, size=10)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.border = Border(bottom=THIN)
            if i in INPUT_COLS:
                c.fill = INPUT_FILL
        if best.get("url"):
            link = ws.cell(row=r, column=12)
            link.hyperlink = best["url"]
            link.font = Font(name=FONT, size=10, color="0563C1", underline="single")
        if g["website"]:
            site = ws.cell(row=r, column=6)
            site.hyperlink = g["website"] if g["website"].startswith("http") else "https://" + g["website"]
            site.font = Font(name=FONT, size=10, color="0563C1", underline="single")
        ws.cell(row=r, column=11).fill = PatternFill("solid", fgColor=PRIORITY_FILL.get(priority, "FFFFFF"))
        ws.cell(row=r, column=14).number_format = "DD.MM.YYYY"
    last = hdr + max(len(rows), 1)
    dv = DataValidation(type="list", formula1=f'"{",".join(STATUSES)}"', allow_blank=True,
                        showErrorMessage=True, errorTitle="Статус", error="Изберете статус от списъка.")
    ws.add_data_validation(dv)
    dv.add(f"M{hdr + 1}:M{last + 200}")
    dv_date = DataValidation(type="date", operator="greaterThan", formula1="DATE(2020,1,1)", allow_blank=True)
    ws.add_data_validation(dv_date)
    dv_date.add(f"N{hdr + 1}:N{last + 200}")
    # Оцветяване на реда по статус
    rng = f"A{hdr + 1}:O{last + 200}"
    for text, color in (("Разговор – интерес", "C6EFCE"), ("Изпратена оферта", "BDD7EE"), ("Среща / оглед", "A9D08E"),
                        ("Не се интересува", "D9D9D9"), ("Грешен номер", "D9D9D9")):
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f'$M{hdr + 1}="{text}"'],
                                                       fill=PatternFill("solid", fgColor=color)))
    ws.freeze_panes = ws.cell(row=hdr + 1, column=3)
    ws.auto_filter.ref = f"A{hdr}:O{last}"
    ws.row_dimensions[hdr].height = 32


def summary(wb: Workbook, sheets: list[str]) -> None:
    ws = wb.create_sheet("Обобщение", 0)
    ws["A1"] = "Фирми с контакт – София и Пловдив"
    ws["A1"].font = Font(name=FONT, bold=True, size=14)
    ws["A2"] = f"Данни към {TODAY.strftime('%d.%m.%Y')}. Попълвайте само жълтите колони: Статус, Дата на обаждане, Бележки."
    ws["A2"].font = Font(name=FONT, italic=True, size=10)
    head = ["Лист", "Фирми", "Висок приоритет", "Звъннати", "Интерес / оферта / среща", "Не се интересуват"]
    for i, h in enumerate(head, 1):
        c = ws.cell(row=4, column=i, value=h)
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
        c.alignment = Alignment(wrap_text=True, horizontal="center")
    for r, s in enumerate(sheets, 5):
        q = f"'{s}'"
        ws.cell(row=r, column=1, value=s)
        ws.cell(row=r, column=2, value=f"=COUNTA({q}!B5:B2000)")
        ws.cell(row=r, column=3, value=f'=COUNTIF({q}!K5:K2000,"Висок*")')
        ws.cell(row=r, column=4, value=f'=COUNTA({q}!B5:B2000)-COUNTIF({q}!M5:M2000,"Не е звъняно")')
        ws.cell(row=r, column=5, value=f'=COUNTIF({q}!M5:M2000,"Разговор – интерес")+COUNTIF({q}!M5:M2000,"Изпратена оферта")'
                                       f'+COUNTIF({q}!M5:M2000,"Среща / оглед")')
        ws.cell(row=r, column=6, value=f'=COUNTIF({q}!M5:M2000,"Не се интересува")')
        for col in range(1, 7):
            ws.cell(row=r, column=col).font = Font(name=FONT)
    total = 5 + len(sheets)
    ws.cell(row=total, column=1, value="Общо").font = Font(name=FONT, bold=True)
    for col in range(2, 7):
        L = get_column_letter(col)
        c = ws.cell(row=total, column=col, value=f"=SUM({L}5:{L}{total - 1})")
        c.font = Font(name=FONT, bold=True)
    notes = [
        ("Как се чете", ""),
        ("Етап сега", "Приблизителен: изчислен от датата на разрешението за строеж и големината на сградата. "
                      "Акт 14 не е публичен. „Потвърдено начало“ = протокол за откриване на строителната площадка."),
        ("Около Акт 14", "Грубият строеж приключва – фасадата се решава и поръчва сега. Най-спешно."),
        ("Груб строеж", "Сградата се издига, фасадата още не е поръчана – най-подходящият момент за оферта."),
        ("Изкоп / начало", "Строежът започва – рано е, но може да се предложи на инвеститора/архитекта за проекта."),
        ("Обновяване", "Енергийно обновяване/саниране – фасадата е част от проекта."),
        ("Обществена поръчка", "Договор за строителство на обществена сграда (училище, градина, читалище …) – "
                               "изпълнителят купува материалите."),
        ("Проектиране", "Договор за проектиране – проектантът избира материалите за фасадата."),
        ("Роля", "Инвеститор = възложител/собственик; Строител = изпълнител; Архитект/Проектант = проектира сградата."),
        ("Един ред = една фирма", "Фирми със същия телефон или имейл (обикновено една група) са на един ред, "
                                  "за да не се звъни два пъти."),
        ("Пример за попълване", "Статус: „Разговор – интерес“ · Дата на обаждане: 08.10.2026 · "
                                "Бележки: „Пратих каталог, да звънна пак 15.10“"),
        ("Източници", "НАГ София и Община Пловдив (регистри на разрешенията за строеж), Търговски регистър, "
                      "КСБ, КАБ, ЦАИС ЕОП (обществени поръчки), сайтове на фирмите."),
    ]
    for i, (k, v) in enumerate(notes, total + 2):
        ws.cell(row=i, column=1, value=k).font = Font(name=FONT, bold=True)
        cell = ws.cell(row=i, column=2, value=v)
        cell.font = Font(name=FONT, size=10)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=i, start_column=2, end_row=i, end_column=6)
        ws.row_dimensions[i].height = 30 if len(v) > 90 else 16
    for col, w in zip("ABCDEF", (24, 14, 14, 12, 16, 16)):
        ws.column_dimensions[col].width = w


def main() -> None:
    sofia_firms, sofia_supervision = sofia()
    plovdiv_firms = plovdiv()
    wb = Workbook()
    wb.remove(wb.active)
    sheet(wb, "София", sofia_firms, "Инвеститори, строители и архитекти на сгради в подходящ момент за фасада "
                                    "(НАГ София, последните 2 години). Подредени по спешност.")
    sheet(wb, "Пловдив", plovdiv_firms, "Инвеститори от регистъра на Община Пловдив и изпълнители/проектанти на "
                                        "обществени сгради в областта (последните 18 месеца). Подредени по спешност.")
    sheet(wb, "София – надзор", sofia_supervision, "Фирми за строителен надзор на тези сгради – знаят строителя и "
                                                   "графика; полезни, когато инвеститорът няма публичен контакт.")
    summary(wb, ["София", "Пловдив", "София – надзор"])
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"София: {len(sofia_firms)} фирми, надзор: {len(sofia_supervision)}, Пловдив: {len(plovdiv_firms)} -> {OUT}")


if __name__ == "__main__":
    main()
